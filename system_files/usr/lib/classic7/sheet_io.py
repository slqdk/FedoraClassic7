"""Opening and saving workbooks for Sheet: Excel 2007+ (.xlsx .xlsm .xltx) through openpyxl,
Excel 97-2003 (.xls) through xlrd, and CSV/text. Charts in .xlsx are read from the file
parts themselves (openpyxl leaves them out) and written back with openpyxl."""
from __future__ import annotations

import colorsys
import csv
import datetime as dt
import io
import os
import posixpath
import zipfile

from lxml import etree

import sheet_formula as F
from sheet_model import (INDEXED_COLORS, Border, Cell, CondRule, PageSetup, Style, Validation, Workbook,
                         normalize_formula)
from sheet_ref import MAX_COLS, MAX_ROWS, cell_name, col_index, parse_cell, parse_range, range_name


class OpenError(Exception):
    pass


# ---- colours -------------------------------------------------------------------------------------
def _tint(hex_color: str, tint: float) -> str:
    r, g, b = (int(hex_color[i:i + 2], 16) / 255 for i in (1, 3, 5))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "#%02x%02x%02x" % (round(r * 255), round(g * 255), round(b * 255))


def color_of(c, theme, default=None):
    """openpyxl Color -> '#rrggbb'."""
    if c is None:
        return default
    try:
        if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) >= 6:
            rgb = c.rgb[-6:]
            if c.rgb.upper() in ("00000000",) and default is not None:
                return default
            out = "#" + rgb.lower()
        elif c.type == "theme" and c.theme is not None:
            idx = int(c.theme)
            # Excel swaps dk1/lt1 and dk2/lt2 in this list
            order = [1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11]
            out = theme[order[idx]] if idx < len(order) else default
        elif c.type == "indexed" and c.indexed is not None:
            i = int(c.indexed)
            if i == 64:
                return default
            out = INDEXED_COLORS[i] if i < len(INDEXED_COLORS) else default
        else:
            return default
        if out and c.tint:
            out = _tint(out, float(c.tint))
        return out
    except (AttributeError, ValueError, TypeError):
        return default


def _border(side, theme) -> Border:
    if side is None or not side.style:
        return Border()
    return Border(side.style, color_of(side.color, theme, "#000000") or "#000000")


def style_from_openpyxl(cell, theme) -> Style:
    f, fill, al, bd = cell.font, cell.fill, cell.alignment, cell.border
    bg = None
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        bg = color_of(fill.fgColor, theme) or color_of(fill.start_color, theme)
    elif fill is not None and getattr(fill, "fill_type", None) not in (None, "none") and hasattr(fill, "fgColor"):
        bg = color_of(fill.fgColor, theme)
    elif fill is not None and getattr(fill, "type", None) == "gradient":
        stops = getattr(fill, "stop", [])
        if stops:
            bg = color_of(stops[0].color, theme)
    return Style(
        font=(f.name or "Calibri") if f else "Calibri",
        size=float(f.sz or 11) if f else 11.0,
        bold=bool(f.b) if f else False,
        italic=bool(f.i) if f else False,
        underline=bool(f.u) and f.u != "none" if f else False,
        strike=bool(f.strike) if f else False,
        color=color_of(f.color, theme, "#000000") if f else "#000000",
        fill=bg,
        halign=(al.horizontal or "general") if al else "general",
        valign=({"centerContinuous": "center", "distributed": "center", "justify": "center"}.get(al.vertical, al.vertical)
                or "bottom") if al else "bottom",
        wrap=bool(al.wrap_text) if al else False,
        indent=int(al.indent or 0) if al else 0,
        rotation=int(al.text_rotation or 0) if al else 0,
        shrink=bool(al.shrink_to_fit) if al else False,
        numfmt=cell.number_format or "General",
        left=_border(bd.left, theme) if bd else Border(),
        right=_border(bd.right, theme) if bd else Border(),
        top=_border(bd.top, theme) if bd else Border(),
        bottom=_border(bd.bottom, theme) if bd else Border(),
        locked=bool(cell.protection.locked) if cell.protection is not None else True,
    )


# ---- xlsx ----------------------------------------------------------------------------------------
def _theme_colors(zf: zipfile.ZipFile) -> list[str]:
    from sheet_model import THEME_COLORS
    try:
        root = etree.fromstring(zf.read("xl/theme/theme1.xml"))
    except (KeyError, etree.XMLSyntaxError):
        return list(THEME_COLORS)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    out = []
    for name in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                 "hlink", "folHlink"):
        el = scheme.find(f"a:{name}", ns) if scheme is not None else None
        val = None
        if el is not None:
            s = el.find("a:srgbClr", ns)
            sc = el.find("a:sysClr", ns)
            val = "#" + s.get("val").lower() if s is not None else (
                "#" + (sc.get("lastClr") or ("000000" if name.startswith("dk") else "ffffff")).lower() if sc is not None else None)
        out.append(val or THEME_COLORS[len(out)])
    return out


def _rels(zf, part):
    d, f = posixpath.split(part)
    path = posixpath.join(d, "_rels", f + ".rels")
    try:
        root = etree.fromstring(zf.read(path))
    except (KeyError, etree.XMLSyntaxError):
        return {}
    out = {}
    for r in root:
        tgt = r.get("Target", "")
        out[r.get("Id")] = (r.get("Type", ""), posixpath.normpath(posixpath.join(d, tgt)) if not tgt.startswith("/") else tgt[1:])
    return out


def _sheet_parts(zf) -> list[tuple[str, str]]:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    root = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = _rels(zf, "xl/workbook.xml")
    out = []
    for s in root.findall("m:sheets/m:sheet", ns):
        rid = s.get("{%s}id" % ns["r"])
        if rid in rels:
            out.append((s.get("name"), rels[rid][1]))
    return out


def _charts_for_sheets(zf) -> dict[str, list]:
    """sheet name -> [Chart] with anchors, read from the drawing parts."""
    import sheet_chart
    xdr = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    out = {}
    try:
        parts = _sheet_parts(zf)
    except (KeyError, etree.XMLSyntaxError):
        return out
    for name, part in parts:
        for rid, (typ, target) in _rels(zf, part).items():
            if not typ.endswith("/drawing"):
                continue
            try:
                droot = etree.fromstring(zf.read(target))
            except (KeyError, etree.XMLSyntaxError):
                continue
            drels = _rels(zf, target)
            for anchor in droot:
                frm = anchor.find(f"{{{xdr}}}from")
                to = anchor.find(f"{{{xdr}}}to")
                gf = anchor.find(f".//{{{sheet_chart.C[1:-1]}}}chart")
                if gf is None:
                    continue
                rid2 = gf.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if rid2 not in drels:
                    continue
                try:
                    ch = sheet_chart.parse_chart(zf.read(drels[rid2][1]))
                except KeyError:
                    ch = None
                if ch is None:
                    continue

                def pos(el):
                    if el is None:
                        return 0, 0
                    return (int(el.findtext(f"{{{xdr}}}row") or 0), int(el.findtext(f"{{{xdr}}}col") or 0))
                fr, fc = pos(frm)
                tr, tc = pos(to) if to is not None else (fr + 15, fc + 8)
                if to is None:
                    ext = anchor.find(f"{{{xdr}}}ext")
                    if ext is not None:
                        tr = fr + max(1, int(int(ext.get("cy", "0")) / 190500))
                        tc = fc + max(1, int(int(ext.get("cx", "0")) / 609600))
                ch.anchor = (fr, fc, max(tr, fr + 1), max(tc, fc + 1))
                out.setdefault(name, []).append(ch)
    return out


def read_xlsx(path: str) -> Workbook:
    try:
        import openpyxl
    except ImportError as e:
        raise OpenError("python3-openpyxl is not installed.") from e
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        raise OpenError("The file is not a valid Excel workbook, or it is damaged.") from e
    with zf:
        theme = _theme_colors(zf)
        charts = _charts_for_sheets(zf)
    try:
        with _no_pivots():
            src = openpyxl.load_workbook(path, data_only=False, keep_links=False, rich_text=False)
    except Exception as e:  # noqa: BLE001 - openpyxl raises many kinds
        if "encrypt" in str(e).lower() or "password" in str(e).lower():
            raise OpenError("This workbook is protected with a password, which Sheet can't open.") from e
        raise OpenError(f"The workbook could not be read: {e}") from e
    cached_book = []

    def cached():
        """Excel's saved results of formulas, read only for workbooks that have formulas."""
        if not cached_book:
            try:
                cached_book.append(openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False))
            except Exception:  # noqa: BLE001
                cached_book.append(None)
        return cached_book[0]
    wb = Workbook()
    wb.theme_colors = theme
    style_cache: dict[int, int] = {}
    for sws in src.worksheets:
        ws = wb.add_sheet(sws.title)
        ws.hidden = sws.sheet_state != "visible"

        def load_cached(title=sws.title):
            book = cached()
            return _cached_values(book, book[title]) if book is not None and title in book.sheetnames else {}
        cvals = _Lazy(load_cached)
        # the cells that exist, not every cell in the stated dimension (A1:IV65536 is 16M cells)
        stored = getattr(sws, "_cells", None)
        rows = ([stored[k] for k in sorted(stored)],) if isinstance(stored, dict) else sws.iter_rows()
        for row in rows:
            for c in row:
                if c.__class__.__name__ == "MergedCell":
                    continue
                v = c.value
                sid = 0
                if c.has_style:
                    key = c.style_id if hasattr(c, "style_id") else id(c._style)
                    sid = style_cache.get(key)
                    if sid is None:
                        sid = wb.style_id(style_from_openpyxl(c, theme))
                        style_cache[key] = sid
                if v is None and not sid:
                    continue
                r, col = c.row - 1, c.column - 1
                if r < 0 or col < 0:
                    continue        # a damaged file without cell positions
                formula = None
                if isinstance(v, str) and v.startswith("=") and c.data_type == "f":
                    formula = normalize_formula(v.replace("_xlfn.", "").replace("_xlws.", ""))
                    v = cvals.get((r, col))
                elif c.data_type == "f" and hasattr(v, "text"):     # array formula
                    formula = "=" + str(getattr(v, "text", "")).lstrip("=")
                    v = cvals.get((r, col))
                if isinstance(v, (dt.datetime, dt.date)):
                    v = F.to_serial(v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day))
                elif isinstance(v, dt.time):
                    v = (v.hour * 3600 + v.minute * 60 + v.second) / 86400
                elif isinstance(v, dt.timedelta):
                    v = v.total_seconds() / 86400
                elif isinstance(v, int) and not isinstance(v, bool):
                    v = float(v)
                elif isinstance(v, str) and v in F.ERRORS:
                    v = F.ERRORS[v]
                ws.cells[(r, col)] = Cell(v, formula, sid)
                if c.comment is not None:
                    ws.comments[(r, col)] = c.comment.text
                if c.hyperlink is not None and c.hyperlink.target:
                    ws.links[(r, col)] = c.hyperlink.target
        # columns and rows
        for key, dim in sws.column_dimensions.items():
            try:
                lo = (dim.min or col_index(key) + 1) - 1
                hi = (dim.max or lo + 1) - 1
            except (TypeError, ValueError):
                continue
            for c in range(lo, min(hi, lo + 300) + 1):
                if dim.width and dim.customWidth is not False:
                    ws.col_widths[c] = float(dim.width)
                if dim.hidden:
                    ws.hidden_cols.add(c)
        for r, dim in sws.row_dimensions.items():
            if dim.height:
                ws.row_heights[r - 1] = float(dim.height)
            if dim.hidden:
                ws.hidden_rows.add(r - 1)
        if sws.sheet_format is not None:
            if sws.sheet_format.defaultColWidth:
                ws.default_width = float(sws.sheet_format.defaultColWidth)
            elif sws.sheet_format.baseColWidth:
                ws.default_width = float(sws.sheet_format.baseColWidth) + 0.71
            if sws.sheet_format.defaultRowHeight:
                ws.default_height = float(sws.sheet_format.defaultRowHeight)
        for m in sws.merged_cells.ranges:
            ws.merges.append((m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1))
        if sws.freeze_panes:
            p = parse_cell(sws.freeze_panes)
            if p and p != (0, 0):
                ws.freeze = p
        if sws.auto_filter is not None and sws.auto_filter.ref:
            rng = parse_range(sws.auto_filter.ref)
            if rng:
                ws.autofilter = rng
        view = sws.sheet_view
        if view is not None:
            ws.show_grid = view.showGridLines is not False
            ws.zoom = int(view.zoomScale or 100)
            if view.showZeros is False:
                ws.show_zeros = False
            try:
                sel = view.selection[0].activeCell if view.selection else None
                p = parse_cell(sel) if sel else None
                if p:
                    ws.active = p
                    ws.selection = [(p[0], p[1], p[0], p[1])]
            except (IndexError, AttributeError):
                pass
        if sws.sheet_properties is not None and sws.sheet_properties.tabColor is not None:
            ws.tab_color = color_of(sws.sheet_properties.tabColor, theme)
        _read_page(sws, ws)
        _read_cond(sws, ws, theme)
        _read_validation(sws, ws)
        ws.charts = charts.get(sws.title, [])
        _read_images(sws, ws)
    # defined names
    try:
        for name, dn in src.defined_names.items():
            if dn.attr_text and not name.startswith("_xlnm"):
                wb.names[name] = "=" + dn.attr_text
    except AttributeError:
        pass
    for sws in src.worksheets:
        try:
            for name, dn in sws.defined_names.items():
                if dn.attr_text and not name.startswith("_xlnm"):
                    wb.local_names[(sws.title.lower(), name.upper())] = "=" + dn.attr_text
        except AttributeError:
            pass
    active = getattr(src.active, "title", None)
    wb.active = next((i for i, ws in enumerate(wb.sheets) if ws.name == active), 0)
    if not wb.sheets:
        wb.add_sheet("Sheet1")
    wb.path = path
    wb.rebuild_dependencies()
    _recalc_keeping_cache(wb)
    wb.modified = False
    return wb


def _recalc_keeping_cache(wb: Workbook):
    """Like Excel opening a file: formulas keep the results saved in the file; only formulas
    without a saved result and volatile ones (TODAY, NOW, RAND...) and what depends on them are
    calculated. Big workbooks open quickly, and functions Sheet lacks keep Excel's value."""
    starts = set(wb._volatile)
    for ws in wb.sheets:
        for (r, c), cell in ws.cells.items():
            if cell.formula is not None and cell.value is None:
                starts.add((ws.name, r, c))
    wb._dirty.clear()
    if starts:
        wb._mark_dirty(starts)
    cache = {k: wb.sheet(k[0]).cells[(k[1], k[2])].value for k in wb._dirty}
    wb.recalc()
    for key, old in cache.items():
        cell = wb.sheet(key[0]).cells.get((key[1], key[2]))
        if cell is not None and cell.value == F.NAME and old is not None:
            cell.value = old


def _read_page(sws, ws):
    p = PageSetup()
    try:
        ps = sws.page_setup
        p.orientation = ps.orientation or "portrait"
        p.paper = {"1": "Letter", "5": "Legal", "8": "A3", "9": "A4", "11": "A5"}.get(str(ps.paperSize or 9), "A4")
        if ps.scale:
            p.scale = int(ps.scale)
        if sws.sheet_properties.pageSetUpPr is not None and sws.sheet_properties.pageSetUpPr.fitToPage:
            p.fit_width = int(ps.fitToWidth if ps.fitToWidth is not None else 1)
            p.fit_height = int(ps.fitToHeight if ps.fitToHeight is not None else 1)
        m = sws.page_margins
        p.margins = (m.left, m.right, m.top, m.bottom)
        p.header_margin, p.footer_margin = m.header, m.footer
        po = sws.print_options
        p.gridlines = bool(po.gridLines)
        p.headings = bool(po.headings)
        p.center_h = bool(po.horizontalCentered)
        p.center_v = bool(po.verticalCentered)
        hf = sws.oddHeader
        p.header = (hf.center.text or "") if hf and hf.center and hf.center.text else ""
        ff = sws.oddFooter
        p.footer = (ff.center.text or "") if ff and ff.center and ff.center.text else ""
        if sws.print_area:
            pa = sws.print_area if isinstance(sws.print_area, str) else sws.print_area[0]
            rng = parse_range(pa.split("!")[-1].split(",")[0])
            if rng:
                p.print_area = rng
        if sws.print_title_rows:
            a, _, b = sws.print_title_rows.replace("$", "").partition(":")
            if a.isdigit():
                p.title_rows = (int(a) - 1, int(b or a) - 1)
    except (AttributeError, TypeError, ValueError):
        pass
    ws.page = p


def _read_cond(sws, ws, theme):
    try:
        cf = sws.conditional_formatting
    except AttributeError:
        return
    for rng_obj in cf:
        ranges = []
        for part in str(rng_obj.sqref).split():
            r = parse_range(part)
            if r:
                ranges.append(r)
        for rule in rng_obj.rules:
            cr = CondRule(ranges=ranges, kind=rule.type or "cellIs", operator=rule.operator or "",
                          formulas=[str(f) for f in (rule.formula or [])], priority=rule.priority or 1,
                          stop=bool(rule.stopIfTrue))
            dxf = rule.dxf
            if dxf is not None:
                if dxf.fill is not None:
                    cr.fill = color_of(getattr(dxf.fill, "bgColor", None), theme) or color_of(
                        getattr(dxf.fill, "fgColor", None), theme)
                if dxf.font is not None:
                    cr.color = color_of(dxf.font.color, theme)
                    cr.bold = bool(dxf.font.b)
            if rule.colorScale is not None:
                cr.colors = [color_of(c, theme, "#ffffff") for c in rule.colorScale.color]
            if rule.dataBar is not None:
                cr.colors = [color_of(rule.dataBar.color, theme, "#638ec6")]
            if rule.type == "top10":
                cr.rank = int(rule.rank or 10)
                cr.percent = bool(rule.percent)
                cr.bottom = bool(rule.bottom)
            if rule.type in ("containsText", "notContainsText", "beginsWith", "endsWith"):
                cr.text = rule.text or ""
            ws.cond.append(cr)


def _read_validation(sws, ws):
    try:
        dvs = sws.data_validations.dataValidation
    except AttributeError:
        return
    for dv in dvs:
        ranges = [r for r in (parse_range(p) for p in str(dv.sqref).split()) if r]
        ws.validations.append(Validation(ranges=ranges, kind=dv.type or "any", formula1=dv.formula1 or "",
                                         formula2=dv.formula2 or "", operator=dv.operator or "between",
                                         allow_blank=bool(dv.allow_blank), show_dropdown=not dv.showDropDown,
                                         error=dv.error or "", prompt=dv.prompt or ""))


def _read_images(sws, ws):
    for img in getattr(sws, "_images", []):
        try:
            data = img._data()
            anchor = img.anchor
            fr = anchor._from
            ws.images.append({"data": data, "anchor": (fr.row, fr.col), "size": (img.width, img.height)})
        except Exception:  # noqa: BLE001 - pictures are optional
            continue


# ---- writing xlsx ------------------------------------------------------------------------------
def _ox_color(hex_color: str | None):
    return ("FF" + hex_color.lstrip("#").upper()) if hex_color else None


def write_xlsx(wb: Workbook, path: str):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.styles import Alignment, Font, PatternFill, Protection, Side
    from openpyxl.styles import Border as OBorder
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    import sheet_chart

    out = openpyxl.Workbook()
    out.remove(out.active)
    cache: dict[int, dict] = {}

    def ox_style(i: int):
        if i in cache:
            return cache[i]
        st = wb.styles[i]

        def side(b: Border):
            return Side(style=b.style or None, color=_ox_color(b.color) if b.style else None)
        d = {
            "font": Font(name=st.font, size=st.size, bold=st.bold, italic=st.italic,
                         underline="single" if st.underline else None, strike=st.strike, color=_ox_color(st.color)),
            "fill": PatternFill("solid", fgColor=_ox_color(st.fill)) if st.fill else PatternFill(fill_type=None),
            "alignment": Alignment(horizontal=None if st.halign == "general" else st.halign,
                                   vertical=None if st.valign == "bottom" else st.valign, wrap_text=st.wrap or None,
                                   indent=st.indent, text_rotation=st.rotation, shrink_to_fit=st.shrink or None),
            "border": OBorder(left=side(st.left), right=side(st.right), top=side(st.top), bottom=side(st.bottom)),
            "number_format": st.numfmt,
            "protection": Protection(locked=st.locked),
        }
        cache[i] = d
        return d

    for ws in wb.sheets:
        o = out.create_sheet(ws.name)
        if ws.hidden:
            o.sheet_state = "hidden"
        for (r, c), cell in sorted(ws.cells.items()):
            oc = o.cell(row=r + 1, column=c + 1)
            if cell.formula is not None:
                oc.value = cell.formula
            else:
                v = cell.value
                if isinstance(v, F.XLError):
                    oc.value = v.code
                    oc.data_type = "e"
                else:
                    oc.value = v
            if cell.style:
                d = ox_style(cell.style)
                oc.font, oc.fill, oc.alignment = d["font"], d["fill"], d["alignment"]
                oc.border, oc.number_format, oc.protection = d["border"], d["number_format"], d["protection"]
            if (r, c) in ws.comments:
                oc.comment = Comment(ws.comments[(r, c)], "Sheet")
            if (r, c) in ws.links:
                oc.hyperlink = ws.links[(r, c)]
        for c, w in ws.col_widths.items():
            o.column_dimensions[get_column_letter(c + 1)].width = w
        for c in ws.hidden_cols:
            o.column_dimensions[get_column_letter(c + 1)].hidden = True
        for r, h in ws.row_heights.items():
            o.row_dimensions[r + 1].height = h
        for r in ws.hidden_rows:
            o.row_dimensions[r + 1].hidden = True
        _merge_all(o, ws.merges)
        if ws.freeze:
            o.freeze_panes = cell_name(*ws.freeze)
        if ws.autofilter:
            o.auto_filter.ref = range_name(*ws.autofilter)
        o.sheet_view.showGridLines = ws.show_grid
        if ws.zoom != 100:
            o.sheet_view.zoomScale = ws.zoom
        if not ws.show_zeros:
            o.sheet_view.showZeros = False
        if ws.tab_color:
            o.sheet_properties.tabColor = ws.tab_color.lstrip("#").upper()
        o.sheet_view.selection[0].activeCell = cell_name(*ws.active)
        o.sheet_view.selection[0].sqref = cell_name(*ws.active)
        _write_page(o, ws)
        _write_cond(o, ws)
        for v in ws.validations:
            dv = DataValidation(type=v.kind if v.kind != "any" else None, formula1=v.formula1 or None,
                                formula2=v.formula2 or None, operator=v.operator if v.kind not in ("list", "custom") else None,
                                allow_blank=v.allow_blank, showDropDown=not v.show_dropdown,
                                error=v.error or None, prompt=v.prompt or None)
            for rng in v.ranges:
                dv.add(range_name(*rng))
            o.add_data_validation(dv)
        for ch in ws.charts:
            try:
                oc_chart = sheet_chart.to_openpyxl(ch)
            except Exception:  # noqa: BLE001 - a chart that can't be written is left out, not the file
                continue
            r1, c1, r2, c2 = ch.anchor
            anchor = TwoCellAnchor()
            anchor._from = AnchorMarker(col=c1, row=r1)
            anchor.to = AnchorMarker(col=c2, row=r2)
            oc_chart.anchor = anchor
            o.add_chart(oc_chart)
        for im in ws.images:
            try:
                from openpyxl.drawing.image import Image as OImage
                img = OImage(io.BytesIO(im["data"]))
                if im.get("size"):
                    img.width, img.height = im["size"]
                o.add_image(img, cell_name(*im["anchor"]))
            except Exception:  # noqa: BLE001
                continue
    for name, text in wb.names.items():
        from openpyxl.workbook.defined_name import DefinedName
        try:
            out.defined_names[name] = DefinedName(name, attr_text=text.lstrip("="))
        except (TypeError, ValueError):
            pass
    out.active = wb.active
    out.properties.creator = os.environ.get("USER", "")
    out.properties.lastModifiedBy = os.environ.get("USER", "")
    tmp = path + ".tmp-classic7"
    out.save(tmp)
    os.replace(tmp, path)


def _write_page(o, ws):
    p = ws.page
    try:
        o.page_setup.orientation = p.orientation
        o.page_setup.paperSize = {"Letter": 1, "Legal": 5, "A3": 8, "A4": 9, "A5": 11}.get(p.paper, 9)
        if p.fit_width or p.fit_height:
            o.sheet_properties.pageSetUpPr.fitToPage = True
            o.page_setup.fitToWidth = p.fit_width
            o.page_setup.fitToHeight = p.fit_height
        elif p.scale != 100:
            o.page_setup.scale = p.scale
        o.page_margins.left, o.page_margins.right, o.page_margins.top, o.page_margins.bottom = p.margins
        o.page_margins.header, o.page_margins.footer = p.header_margin, p.footer_margin
        o.print_options.gridLines = p.gridlines
        o.print_options.headings = p.headings
        o.print_options.horizontalCentered = p.center_h
        o.print_options.verticalCentered = p.center_v
        if p.header:
            o.oddHeader.center.text = p.header
        if p.footer:
            o.oddFooter.center.text = p.footer
        if p.print_area:
            o.print_area = range_name(*p.print_area)
        if p.title_rows:
            o.print_title_rows = f"{p.title_rows[0] + 1}:{p.title_rows[1] + 1}"
    except (AttributeError, TypeError, ValueError):
        pass


def _write_cond(o, ws):
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle
    for cr in ws.cond:
        rng = " ".join(range_name(*r) for r in cr.ranges)
        try:
            fill = PatternFill(bgColor=_ox_color(cr.fill), fill_type="solid") if cr.fill else None
            font = Font(color=_ox_color(cr.color), bold=cr.bold or None) if (cr.color or cr.bold) else None
            if cr.kind == "cellIs":
                rule = CellIsRule(operator=cr.operator, formula=cr.formulas, fill=fill, font=font)
            elif cr.kind == "expression":
                rule = FormulaRule(formula=cr.formulas, fill=fill, font=font)
            elif cr.kind == "colorScale" and len(cr.colors) >= 2:
                if len(cr.colors) == 2:
                    rule = ColorScaleRule(start_type="min", start_color=_ox_color(cr.colors[0]),
                                          end_type="max", end_color=_ox_color(cr.colors[1]))
                else:
                    rule = ColorScaleRule(start_type="min", start_color=_ox_color(cr.colors[0]),
                                          mid_type="percentile", mid_value=50, mid_color=_ox_color(cr.colors[1]),
                                          end_type="max", end_color=_ox_color(cr.colors[2]))
            elif cr.kind == "dataBar":
                rule = DataBarRule(start_type="min", end_type="max", color=_ox_color(cr.colors[0] if cr.colors else "#638ec6"))
            else:
                dxf = DifferentialStyle(fill=fill, font=font)
                rule = Rule(type=cr.kind, dxf=dxf, rank=cr.rank if cr.kind == "top10" else None,
                            percent=cr.percent or None, bottom=cr.bottom or None, text=cr.text or None,
                            formula=cr.formulas or None)
            o.conditional_formatting.add(rng, rule)
        except (TypeError, ValueError):
            continue


# ---- xls (Excel 97-2003) -------------------------------------------------------------------------
def read_xls(path: str) -> Workbook:
    try:
        import xlrd
    except ImportError as e:
        raise OpenError("python3-xlrd is not installed, so old .xls files can't be opened.") from e
    try:
        book = xlrd.open_workbook(path, formatting_info=True, ragged_rows=True)
    except NotImplementedError:
        book = xlrd.open_workbook(path, ragged_rows=True)
    except Exception as e:  # noqa: BLE001
        msg = str(e)
        if "encrypt" in msg.lower() or "password" in msg.lower():
            raise OpenError("This workbook is protected with a password, which Sheet can't open.") from e
        raise OpenError(f"The workbook could not be read: {msg}") from e
    wb = Workbook()
    palette = getattr(book, "colour_map", {}) or {}

    def col(idx):
        rgb = palette.get(idx) if idx is not None else None
        return "#%02x%02x%02x" % rgb if rgb else None

    borders = {0: "", 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair"}
    halign = {0: "general", 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous"}
    valign = {0: "top", 1: "center", 2: "bottom", 3: "center", 4: "center"}
    xf_styles: dict[int, int] = {}

    def style_for(xf_index: int) -> int:
        if xf_index in xf_styles:
            return xf_styles[xf_index]
        try:
            xf = book.xf_list[xf_index]
        except (AttributeError, IndexError):
            return 0
        font = book.font_list[xf.font_index]
        fmt = book.format_map.get(xf.format_key)
        bd = xf.border
        bg = None
        if xf.background.fill_pattern:
            bg = col(xf.background.pattern_colour_index)
        st = Style(font=font.name or "Arial", size=font.height / 20 if font.height else 10, bold=bool(font.bold),
                   italic=bool(font.italic), underline=bool(font.underline_type), strike=bool(font.struck_out),
                   color=col(font.colour_index) or "#000000", fill=bg,
                   halign=halign.get(xf.alignment.hor_align, "general"),
                   valign=valign.get(xf.alignment.vert_align, "bottom"), wrap=bool(xf.alignment.text_wrapped),
                   indent=xf.alignment.indent_level, rotation=xf.alignment.rotation if xf.alignment.rotation <= 180 else 0,
                   numfmt=(fmt.format_str if fmt else "General") or "General",
                   left=Border(borders.get(bd.left_line_style, ""), col(bd.left_colour_index) or "#000000"),
                   right=Border(borders.get(bd.right_line_style, ""), col(bd.right_colour_index) or "#000000"),
                   top=Border(borders.get(bd.top_line_style, ""), col(bd.top_colour_index) or "#000000"),
                   bottom=Border(borders.get(bd.bottom_line_style, ""), col(bd.bottom_colour_index) or "#000000"))
        if st.numfmt == "General" and fmt is None:
            st = Style(**{**st.__dict__, "numfmt": "General"})
        sid = wb.style_id(st)
        xf_styles[xf_index] = sid
        return sid

    for sh in book.sheets():
        ws = wb.add_sheet(sh.name)
        ws.hidden = getattr(sh, "visibility", 0) != 0
        types, values = sh._cell_types, sh._cell_values
        xfs = getattr(sh, "_cell_xf_indexes", None)
        visible_blank: dict[int, bool] = {}
        for r in range(sh.nrows):
            trow, vrow = types[r], values[r]
            xrow = xfs[r] if xfs else None
            for c in range(len(trow)):
                t, v = trow[c], vrow[c]
                xf = xrow[c] if xrow is not None and c < len(xrow) else None
                try:
                    sid = style_for(xf) if xf is not None and xf >= 0 else 0
                except Exception:  # noqa: BLE001
                    sid = 0
                if t == xlrd.XL_CELL_EMPTY or t == xlrd.XL_CELL_BLANK:
                    # a blank cell matters only if it shows something (fill or border)
                    if sid:
                        show = visible_blank.get(sid)
                        if show is None:
                            st = wb.styles[sid]
                            show = bool(st.fill or st.left.style or st.right.style or st.top.style or st.bottom.style)
                            visible_blank[sid] = show
                        if show:
                            ws.cells[(r, c)] = Cell(None, None, sid)
                    continue
                if t == xlrd.XL_CELL_NUMBER or t == xlrd.XL_CELL_DATE:
                    v = float(v)
                elif t == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif t == xlrd.XL_CELL_ERROR:
                    v = F.ERRORS.get(xlrd.error_text_from_code.get(v, "#VALUE!"), F.VALUE)
                ws.cells[(r, c)] = Cell(v, None, sid)
        try:
            for (r1, r2, c1, c2) in sh.merged_cells:
                ws.merges.append((r1, c1, r2 - 1, c2 - 1))
            for c, info in sh.colinfo_map.items():
                ws.col_widths[c] = info.width / 256
                if info.hidden:
                    ws.hidden_cols.add(c)
            for r, info in sh.rowinfo_map.items():
                if info.height:
                    ws.row_heights[r] = info.height / 20
                if info.hidden:
                    ws.hidden_rows.add(r)
            if sh.horz_split_pos or sh.vert_split_pos:
                ws.freeze = (sh.horz_split_pos, sh.vert_split_pos)
        except AttributeError:
            pass
    if not wb.sheets:
        wb.add_sheet("Sheet1")
    wb.path = path
    wb.rebuild_dependencies()
    wb.modified = False
    return wb


# ---- csv -----------------------------------------------------------------------------------------
def read_csv(path: str) -> Workbook:
    import sheet_numfmt as NF
    with open(path, "rb") as f:
        raw = f.read()
    if b"\x00" in raw[:4096]:
        raise OpenError("This file is in a format Sheet can't open (perhaps a very old Excel version).")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    wb = Workbook()
    ws = wb.add_sheet(os.path.splitext(os.path.basename(path))[0][:31] or "Sheet1")
    for r, row in enumerate(csv.reader(io.StringIO(text, newline=""), dialect)):
        for c, val in enumerate(row):
            if val == "":
                continue
            if val.startswith("="):
                try:
                    F.parse(val)
                    ws.cells[(r, c)] = Cell(None, val, 0)
                    continue
                except F.FormulaError:
                    pass
            v, fmt = NF.parse_input(val)
            sid = wb.style_id(Style(numfmt=fmt)) if fmt else 0
            ws.cells[(r, c)] = Cell(v, None, sid)
    wb.path = path
    wb.rebuild_dependencies()
    wb.recalc("all")
    wb.modified = False
    return wb


def write_csv(wb: Workbook, path: str, sep: str = ","):
    ws = wb.ws
    used = ws.used_range()
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=sep)
        if not used:
            return
        r1, c1, r2, c2 = 0, 0, used[2], used[3]
        for r in range(r1, r2 + 1):
            w.writerow([wb.display(ws, r, c)[0] for c in range(c1, c2 + 1)])


# ---- entry points ---------------------------------------------------------------------------------
class _no_pivots:
    """Sheet has no pivot tables, and openpyxl parsing their caches can take minutes."""

    def __enter__(self):
        self.saved = []
        try:
            import openpyxl.reader.excel as ex
            import openpyxl.reader.workbook as wbr
            self.saved = [(ex.TableDefinition, "rel_type", ex.TableDefinition.rel_type),
                          (wbr.WorkbookParser, "pivot_caches", wbr.WorkbookParser.__dict__["pivot_caches"])]
            ex.TableDefinition.rel_type = "urn:classic7:no-pivot-tables"
            wbr.WorkbookParser.pivot_caches = property(lambda self: {})
        except (ImportError, AttributeError, KeyError):
            pass
        return self

    def __exit__(self, *exc):
        for obj, name, value in self.saved:
            setattr(obj, name, value)
        return False


class _Lazy:
    def __init__(self, load):
        self.load, self.data = load, None

    def get(self, key):
        if self.data is None:
            self.data = self.load()
        return self.data.get(key)


def _cached_values(book, sheet) -> dict:
    """The values Excel saved for formula cells. Only the stored cells are parsed (iter_rows
    would pad every missing row and column of the stated dimension)."""
    out = {}
    try:
        from openpyxl.worksheet._reader import WorkSheetParser
        with sheet._get_source() as src:
            parser = WorkSheetParser(src, sheet._shared_strings, data_only=True, epoch=book.epoch,
                                     date_formats=book._date_formats, timedelta_formats=book._timedelta_formats)
            for _idx, row in parser.parse():
                for c in row:
                    if c.get("value") is not None:
                        out[(c["row"] - 1, c["column"] - 1)] = c["value"]
        return out
    except Exception:  # noqa: BLE001 - fall back to the public, slower way
        out.clear()
    try:
        for row in sheet.iter_rows():
            for c in row:
                if getattr(c, "value", None) is not None and hasattr(c, "row"):
                    out[(c.row - 1, c.column - 1)] = c.value
    except Exception:  # noqa: BLE001 - cached values are a nicety
        pass
    return out


def _merge_all(o, merges):
    """openpyxl's merge_cells checks each new range against all earlier ones (slow with
    thousands); our merges never overlap, so add them to the set directly."""
    try:
        from openpyxl.worksheet.merge import MergedCellRange
        ranges = o.merged_cells.ranges
        if not isinstance(ranges, set):
            raise TypeError
        for m in merges:
            mcr = MergedCellRange(o, range_name(*m))
            ranges.add(mcr)
            o._clean_merge_range(mcr)
    except (ImportError, TypeError, AttributeError):
        for m in merges:
            o.merge_cells(start_row=m[0] + 1, start_column=m[1] + 1, end_row=m[2] + 1, end_column=m[3] + 1)


def open_any(path: str) -> Workbook:
    try:
        return _open_any(path)
    except (OpenError, OSError, MemoryError):
        raise
    except Exception as e:  # noqa: BLE001 - damaged files break the readers in many ways
        raise OpenError(f"The workbook is damaged and could not be read ({type(e).__name__}: {e}).") from e


def _open_any(path: str) -> Workbook:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsb":
        raise OpenError("Excel Binary Workbooks (.xlsb) can't be opened. In Excel, save the workbook as an "
                        "Excel Workbook (.xlsx).")
    with open(path, "rb") as f:
        head = f.read(8)
    if head.startswith(b"PK"):
        return read_xlsx(path)
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return read_xls(path)
    if ext in (".csv", ".txt", ".tsv", ".prn"):
        return read_csv(path)
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        raise OpenError("The file is not a valid Excel workbook, or it is damaged.")
    return read_csv(path)


def save_any(wb: Workbook, path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt"):
        write_csv(wb, path, "," if ext == ".csv" else "\t")
    else:
        write_xlsx(wb, path)


LIMITS = (MAX_ROWS, MAX_COLS)
